"""
ORDO 학사 기준정보 시드 생성기
------------------------------------------------------------
입력 (예소 폴더):
  - 경희대_국제캠_2026_졸업관리데이터.xlsx          (교양구조 · 졸업요건 · 교과목마스터)
  - 경희대_교육과정기본구조/{YYYY}학번/*.xlsx        (학번별 전공 이수학점, '국제' 시트)
출력:
  - src/main/resources/db/migration/V2__seed_catalog.sql

실행:
  pip install openpyxl
  python tools/seed/generate_seed.py --data-dir "<예소 폴더 경로>" --out src/main/resources/db/migration/V2__seed_catalog.sql

엑셀이 수정되면 이 스크립트를 다시 돌려 V2 를 재생성한다.
(이미 V2 가 적용된 DB 라면 새 번호(V2xx)로 만들거나, 로컬 DB 를 초기화한 뒤 적용)
"""
import argparse
import os
import re
import sys
from collections import OrderedDict

import openpyxl

YEARS = range(2020, 2027)

# 단과대 이름 표기 통일
COLLEGE_ALIAS = {"예술디자인대학": "예술·디자인대학"}

# 이수구분(한글) -> DB/Java enum 값
CLASSIFICATION = {
    "전공기초": "MAJOR_BASIC",
    "전공필수": "MAJOR_REQUIRED",
    "전공선택": "MAJOR_ELECTIVE",
    "전공선택(교직)": "MAJOR_ELECTIVE",
    "교직": "TEACHING",
    "교직전선": "TEACHING_MAJOR",
}

# 학번별 전공(표시명) -> 2026 교과목마스터의 '학과/학부' 단위명
# 매핑이 없으면 표시명 그대로 사용. None 이면 교과목 매칭 불가(수동 입력).
UNIT_ALIAS = {
    "기계공학과": "기계공학부",
    "기계공학부 기계공학": "기계공학부",
    "기계공학부 지능로봇공학": "기계공학부",
    "기계공학부 항공우주모빌리티": "기계공학부",
    "정보전자신소재공학과": "신소재공학과",
    "환경학및환경공학과 환경학": "환경학전공",
    "환경학및환경공학과 환경학과": "환경학전공",
    "환경학및환경공학과 환경공학": "환경공학전공",
    "환경학및환경공학과 환경공학과": "환경공학전공",
    "식물환경신소재공학과": "식물·환경신소재공학과",
    "유전공학과": "유전생명학과",
    "유전생명공학과": "유전생명학과",
    "한방재료공학과": "한방생명공학과",
    "컴퓨터공학과": "컴퓨터공학부 컴퓨터공학과",
    "전자공학과": "전자공학부 전자공학과",
    "반도체공학과": "전자공학부 반도체공학과",
    "전자정보공학부 전자공학과": "전자공학부 전자공학과",
    "전자정보공학부 반도체공학과": "전자공학부 반도체공학과",
    "한국어학과": "한국어학과 한국어학전공",
    "글로벌한국학과": None,
    # 융합전공
    "웨어러블디자인테크놀로지 융합전공": "웨어러블 디자인 테크놀로지 융합전공",
    "첨단바이오신소재 융합전공": "첨단바이오소재 융합전공",
    "스마트팜공학 융합전공": None,
}


def q(v):
    """SQL 리터럴"""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    s = str(v).replace("\\", "\\\\").replace("'", "''")
    return f"'{s}'"


def num(v):
    """'-', None, '' -> None / 숫자 -> int"""
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "-"):
        return None
    m = re.match(r"^\d+", s)
    return int(m.group()) if m else None


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def insert_batches(table, cols, rows, size=200):
    out = []
    for i in range(0, len(rows), size):
        chunk = rows[i:i + size]
        values = ",\n".join("(" + ", ".join(q(v) for v in r) + ")" for r in chunk)
        out.append(f"INSERT INTO {table} ({', '.join(cols)}) VALUES\n{values};\n")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    master_wb = openpyxl.load_workbook(
        os.path.join(a.data_dir, "경희대_국제캠_2026_졸업관리데이터.xlsx"), read_only=True, data_only=True)

    # ---------- 1) 교과목 마스터 ----------
    master_rows = [list(r) for r in master_wb["⑤교과목마스터"].iter_rows(values_only=True)][3:]
    courses = OrderedDict()   # (unit, code, cls) -> row
    units = set()
    for r in master_rows:
        college, unit, cls_ko, name, code, credits, grade, sem, track = (r + [None] * 9)[:9]
        if not cls_ko or not name:
            continue
        unit = txt(unit) or txt(college)          # 자유전공학부는 학과칸이 비어 있음
        cls = CLASSIFICATION[cls_ko.strip()]
        credit_txt = txt(credits) or "0"
        variable = "-" in credit_txt
        credit_val = max(int(x) for x in re.findall(r"\d+", credit_txt)) if re.findall(r"\d+", credit_txt) else 0
        key = (unit, txt(code), cls)
        units.add(unit)
        if key in courses:
            # 같은 과목이 트랙별로 중복 기재된 경우 트랙만 합침
            old = courses[key]
            if track and track not in (old[9] or ""):
                old[9] = f"{old[9]}, {track}" if old[9] else track
            continue
        courses[key] = [COLLEGE_ALIAS.get(college, college), unit, txt(code), name.strip(), cls,
                        credit_val, variable, txt(grade), txt(sem), txt(track)]

    # ---------- 2) 학번별 전공 이수학점 (국제캠) ----------
    colleges = OrderedDict()
    majors = OrderedDict()     # display_name -> dict
    reqs = []
    for y in YEARS:
        path = os.path.join(a.data_dir, "경희대_교육과정기본구조", f"{y}학번",
                            f"경희대_전공별교육과정기본구조_{y}학번.xlsx")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["국제"]
        c = d = None
        for r in [list(x) for x in ws.iter_rows(values_only=True)][2:]:
            r = (r + [None] * 17)[:17]
            if not any(v is not None for v in r):
                continue
            if r[0]:
                c = COLLEGE_ALIAS.get(r[0].strip(), r[0].strip())
            if r[1]:
                d = r[1].strip()
            major_name = txt(r[2])
            convergence = (c == "융합전공")
            if convergence:
                display = f"{d} 융합전공"
            else:
                display = f"{d} {major_name}" if major_name else d
            colleges.setdefault(c, None)
            if display not in majors:
                unit = UNIT_ALIAS.get(display, display)
                if unit is not None and unit not in units:
                    unit = None
                majors[display] = dict(college=c, department=d, name=major_name,
                                       display=display, unit=unit, convergence=convergence)
            reqs.append(dict(display=display, year=y, total=num(r[3]),
                             s=[num(v) for v in r[4:9]], dbl=[num(v) for v in r[9:14]],
                             minor=[num(v) for v in r[14:17]]))

    # ---------- 3) 2026 졸업요건 부가항목 (SW/영어강의/논문/TOPIK/인증) ----------
    extras = {}
    for r in [list(x) for x in master_wb["④학과별_졸업요건"].iter_rows(values_only=True)][3:]:
        r = (r + [None] * 8)[:8]
        if not r[1]:
            continue
        extras[r[1].strip()] = dict(sw=txt(r[3]), eng=txt(r[4]), thesis=txt(r[5]),
                                    topik=txt(r[6]), cert=txt(r[7]))

    def extra_for(m):
        for k in (m["unit"], m["display"]):
            if k and k in extras:
                return extras[k]
        return None

    # ---------- 4) 교양 (후마니타스칼리지 2026) ----------
    gen_rows = [list(x) for x in master_wb["②교양_기본구조"].iter_rows(values_only=True)][3:]
    gen_required = []
    for r in gen_rows:
        r = (r + [None] * 6)[:6]
        if r[0] == "필수교과" and r[2] and not str(r[2]).startswith("("):
            gen_required.append([2026, txt(r[1]), r[2].strip(), num(r[3]), txt(r[4]), txt(r[5])])

    # ---------- SQL 출력 ----------
    college_ids = {n: i + 1 for i, n in enumerate(colleges)}
    major_ids = {n: i + 1 for i, n in enumerate(majors)}
    out = ["-- ============================================================",
           "-- V2__seed_catalog.sql  (tools/seed/generate_seed.py 로 자동 생성 — 직접 수정 금지)",
           "-- 출처: 경희대 국제캠 2026 교육과정 / 전공별교육과정기본구조 2020~2026학번",
           "-- ============================================================\n"]
    out.append(insert_batches("colleges", ["id", "name"],
                              [[college_ids[n], n] for n in colleges]))
    out.append(insert_batches(
        "majors", ["id", "college_id", "department_name", "major_name", "display_name", "course_unit", "convergence"],
        [[major_ids[k], college_ids[m["college"]], m["department"], m["name"], m["display"], m["unit"], m["convergence"]]
         for k, m in majors.items()]))

    req_rows = []
    for rq in reqs:
        m = majors[rq["display"]]
        ex = extra_for(m) if rq["year"] == 2026 else None
        req_rows.append([major_ids[rq["display"]], rq["year"], rq["total"], *rq["s"], *rq["dbl"], *rq["minor"],
                         ex and ex["sw"], ex and ex["eng"], ex and ex["thesis"], ex and ex["topik"], ex and ex["cert"]])
    out.append(insert_batches("graduation_requirements", [
        "major_id", "admission_year", "total_credits",
        "basic_credits", "required_credits", "elective_credits", "major_total_credits", "other_major_credits",
        "double_basic_credits", "double_required_credits", "double_elective_credits", "double_total_credits",
        "double_other_major_credits",
        "minor_required_credits", "minor_elective_credits", "minor_total_credits",
        "sw_requirement", "english_lecture_requirement", "thesis_requirement", "topik_requirement",
        "competency_certification"], req_rows))

    out.append(insert_batches("general_education_requirements", [
        "admission_year", "required_credits", "distribution_credits", "distribution_min_areas",
        "free_credits", "total_credits"], [[2026, 17, 9, 3, 3, 29]]))
    out.append(insert_batches("general_education_required_courses", [
        "admission_year", "group_name", "course_name", "credits", "recommended_grade", "note"], gen_required))

    out.append(insert_batches("courses", [
        "college_name", "unit_name", "course_code", "name", "classification",
        "credits", "variable_credits", "target_grade", "open_semester", "track"], list(courses.values())))

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as f:
        f.write("\n".join(out))

    unmapped = [m["display"] for m in majors.values() if m["unit"] is None]
    print(f"colleges={len(colleges)} majors={len(majors)} requirements={len(req_rows)} "
          f"gen_required={len(gen_required)} courses={len(courses)}", file=sys.stderr)
    print(f"교과목 매칭 불가 전공(course_unit=NULL): {unmapped}", file=sys.stderr)


if __name__ == "__main__":
    main()
